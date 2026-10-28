"""DOCX text extraction via python-docx.

A .docx has no page concept until it is laid out, so paragraphs are grouped
into synthetic "pages" purely to keep the evidence-reference contract uniform
with PDFs. The grouping is explicit about being an approximation.
"""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.errors import MalformedDocument
from app.schemas.document import PageText

# Roughly a page of body text. Only used to chunk paragraphs for references.
CHARS_PER_SYNTHETIC_PAGE = 1800


def extract_docx(data: bytes) -> list[PageText]:
    try:
        document = Document(BytesIO(data))
    except PackageNotFoundError as exc:
        raise MalformedDocument(
            "That file could not be opened as a Word document."
        ) from exc

    blocks: list[str] = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    pages: list[PageText] = []
    buffer: list[str] = []
    size = 0
    for block in blocks:
        buffer.append(block)
        size += len(block)
        if size >= CHARS_PER_SYNTHETIC_PAGE:
            pages.append(PageText(page_number=len(pages) + 1, text="\n".join(buffer)))
            buffer, size = [], 0
    if buffer:
        pages.append(PageText(page_number=len(pages) + 1, text="\n".join(buffer)))

    return pages
